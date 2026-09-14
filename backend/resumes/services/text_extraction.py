"""Extract plain text from uploaded resume files (PDF / DOCX)."""
import io


def extract_text_from_pdf(file_obj) -> str:
    import pdfplumber

    text_parts = []
    file_obj.seek(0)
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ''
            text_parts.append(page_text)
    return '\n'.join(text_parts).strip()


def extract_text_from_docx(file_obj) -> str:
    import docx

    file_obj.seek(0)
    document = docx.Document(io.BytesIO(file_obj.read()))
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(' '.join(cell.text for cell in row.cells))
    return '\n'.join(paragraphs).strip()


def extract_text(file_obj, file_type: str) -> str:
    if file_type == 'pdf':
        return extract_text_from_pdf(file_obj)
    if file_type == 'docx':
        return extract_text_from_docx(file_obj)
    raise ValueError(f'Unsupported file type: {file_type}')


def infer_file_type(filename: str) -> str:
    lower = filename.lower()
    if lower.endswith('.pdf'):
        return 'pdf'
    if lower.endswith('.docx'):
        return 'docx'
    raise ValueError('Only .pdf and .docx resumes are supported.')
